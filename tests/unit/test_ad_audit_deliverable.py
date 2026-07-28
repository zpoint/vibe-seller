"""The audit is graded on what the run DELIVERED, over every combo the
store declares, with a per-target row behind every targeting table.

Three holes, all observed on live runs of the same weekly audit:

* the gate graded the agent's chat summary while the finished report sat
  in the workspace (24 denials, then a fail-open that stored the
  summary);
* the agent wrote the combo list itself, so a store configured for five
  marketplaces was audited for one and passed;
* a targeting "table" of one 整体活动 aggregate row counted as a drill,
  because the presence check was satisfied by the header row alone.
"""

import json
import os

from openpyxl import Workbook
import pytest

from app.ai.stop_gates import (
    ad_completeness_review as acr,
    ad_scope,
    ad_scope as sc,
)
from app.routers.tasks_files import resolve_audit_deliverable

# One drilled campaign: per-keyword rows + a same-window 对账 line.
_DRILLED = (
    '### 600000000001 | acme widgets 004 manual | SP\n'
    '| 关键词 | 出价 | 点击 | 花费 | ROAS | 建议 |\n'
    '|---|---|---|---|---|---|\n'
    '| widget | 1.20 | 6 | 5.00 | 3.5 | 维持 |\n'
    '| widget large | 1.00 | 4 | 3.00 | 4.1 | 提高至 1.30 |\n'
    '| **合计** | — | 10 | 8.00 | 3.7 | — |\n'
    '搜索词对账: 定向花费 USD 8.00 / 点击 10 = '
    '搜索词花费 USD 8.00 / 点击 10 (✓)\n'
)
# Same campaign, but the targeting table restates the campaign total and
# names no keyword — the shape that used to pass.
_AGGREGATE_ONLY = (
    '### 600000000001 | acme widgets 004 manual | SP\n'
    '| 定位层汇总 | 出价 | 点击 | 花费 | ROAS | 建议 |\n'
    '|---|---|---|---|---|---|\n'
    '| 整体活动 | — | 10 | 8.00 | 3.7 | 钻取关键词级后再降价 |\n'
    '搜索词对账: 定向花费 USD 8.00 / 点击 10 = '
    '搜索词花费 USD 8.00 / 点击 10 (✓)\n'
)
_SUMMARY_SECTION = (
    '## 汇总建议\n'
    '各 combo 花费/销售/ROAS 总览与按影响排序的行动清单：本次覆盖 '
    'Amazon SA，ROAS 3.5，建议维持核心词出价并持续监控搜索词转化。\n'
)


def _report(body: str) -> str:
    return (
        '## Amazon SA\n**进度**: drilled 1/1 active (1 total, 1 pages)\n' + body
    ) + _SUMMARY_SECTION


def _setup(monkeypatch, tmp_path, task_id, *, scope=None, targets=None):
    monkeypatch.setattr(ad_scope, 'VIBE_SELLER_DIR', tmp_path)
    tdir = tmp_path / 'tasks' / task_id
    tdir.mkdir(parents=True, exist_ok=True)
    if scope is not None:
        (tdir / 'AUDIT_SCOPE.json').write_text(
            json.dumps({'combos': scope}), encoding='utf-8'
        )
    if targets is not None:
        ad_scope.write_declared_targets(tdir, targets)
    acr.reset_progress(task_id)
    return tdir


def _gaps(text, task_id):
    deny = acr.check(text, task_id=task_id, track=False)
    return list(deny.gaps) if deny else []


_SA_SCOPE = [
    {
        'platform': 'amazon',
        'country': 'SA',
        'active_ids': ['600000000001'],
        'total_active': 1,
        'total_active_source': 'chip:Live 1',
    }
]


@pytest.mark.unit
class TestDeliverableOverNarration:
    """A run is graded on the report it wrote, not the prose about it."""

    def test_narration_resolves_to_the_report(self, tmp_path):
        (tmp_path / 'AD_AUDIT_2026-01-02.md').write_text('## Amazon SA\n')
        found = resolve_audit_deliverable(
            tmp_path, 'Amazon SA 审计完成 — drilled 20/20，见报告文件。'
        )
        assert found is not None
        assert found.name == 'AD_AUDIT_2026-01-02.md'

    def test_inlined_report_is_never_downgraded(self, tmp_path):
        # A stale file from an earlier round must not replace a full
        # report the agent submitted inline.
        (tmp_path / 'AD_AUDIT_2026-01-01.md').write_text('stale')
        assert resolve_audit_deliverable(tmp_path, _report(_DRILLED)) is None

    def test_no_report_file_leaves_submission_alone(self, tmp_path):
        assert resolve_audit_deliverable(tmp_path, 'done, 20/20') is None

    def test_execution_task_keeps_its_own_result(self, tmp_path):
        # Phase-4 execution reads a prior audit; that audit is its INPUT.
        (tmp_path / 'AD_AUDIT_2026-01-02.md').write_text('## Amazon SA\n')
        (tmp_path / 'EXECUTION_LOG.md').write_text('applied 12 bid changes')
        assert (
            resolve_audit_deliverable(tmp_path, '执行完毕，12 项已应用') is None
        )

    def test_newest_report_wins(self, tmp_path):
        old = tmp_path / 'AD_AUDIT_2026-01-01.md'
        new = tmp_path / 'AD_AUDIT_2026-01-02.md'
        old.write_text('old')
        new.write_text('new')
        os.utime(old, (1_600_000_000, 1_600_000_000))
        os.utime(new, (1_700_000_000, 1_700_000_000))
        assert resolve_audit_deliverable(tmp_path, 'done').name == new.name


@pytest.mark.unit
class TestDeclaredComboCoverage:
    """The store's marketplaces are the server's fact, not the agent's."""

    def test_omitted_combo_is_a_gap(self, monkeypatch, tmp_path):
        _setup(
            monkeypatch,
            tmp_path,
            't-decl',
            scope=_SA_SCOPE,
            targets={'amazon': ['SA', 'AE'], 'noon': ['SA']},
        )
        gaps = _gaps(_report(_DRILLED), 't-decl')
        assert any('amazon AE' in g for g in gaps), gaps
        assert any('noon SA' in g for g in gaps), gaps
        # The combo that WAS audited raises nothing.
        assert not any('amazon SA」根本没进' in g for g in gaps), gaps

    def test_combo_may_be_declared_empty(self, monkeypatch, tmp_path):
        # "No live campaigns in this marketplace" is a finding, not an
        # omission — it just has to be written down.
        _setup(
            monkeypatch,
            tmp_path,
            't-empty',
            scope=[
                *_SA_SCOPE,
                {
                    'platform': 'amazon',
                    'country': 'AE',
                    'active_ids': [],
                    'total_active': 0,
                    'total_active_source': 'chip:Live 0',
                },
            ],
            targets={'amazon': ['SA', 'AE']},
        )
        gaps = _gaps(_report(_DRILLED), 't-empty')
        # No [基线] gap AT ALL. Asserting only that the missing-combo gap
        # is gone would miss a deadlock: the empty-active_ids branch used
        # to answer "delete the entry", and deleting it re-raises the
        # missing-combo gap — two remedies, no legal move.
        assert not any('[基线]' in g for g in gaps), gaps

    def test_empty_combo_still_needs_an_observed_zero(
        self, monkeypatch, tmp_path
    ):
        # active_ids: [] with no total_active is a TRUNCATED enumeration
        # wearing the same clothes as an empty marketplace — still a gap.
        _setup(
            monkeypatch,
            tmp_path,
            't-empty-untotalled',
            scope=[
                *_SA_SCOPE,
                {'platform': 'amazon', 'country': 'AE', 'active_ids': []},
            ],
            targets={'amazon': ['SA', 'AE']},
        )
        gaps = _gaps(_report(_DRILLED), 't-empty-untotalled')
        assert any('active_ids 是空的' in g for g in gaps), gaps
        # …and the remedy must NOT be "delete the entry" for a combo the
        # server declared — that is the deadlock.
        assert not any('整条删掉' in g for g in gaps), gaps

    def test_empty_combo_may_not_also_hedge(self, monkeypatch, tmp_path):
        # ``exhaustive: false`` + ``active_ids: []`` reads "I audited a
        # subset, and the subset was nothing" — zero obligations wearing
        # the narrow-task escape hatch. Seen live on two Amazon markets.
        _setup(
            monkeypatch,
            tmp_path,
            't-hedge',
            scope=[
                *_SA_SCOPE,
                {
                    'platform': 'amazon',
                    'country': 'AE',
                    'active_ids': [],
                    'total_active': 0,
                    'total_active_source': 'chip:Live 0',
                    'exhaustive': False,
                },
            ],
            targets={'amazon': ['SA', 'AE']},
        )
        gaps = _gaps(_report(_DRILLED), 't-hedge')
        assert any('exhaustive' in g and 'amazon AE' in g for g in gaps), gaps

    def test_empty_combo_contradicting_history_is_challenged(
        self, monkeypatch, tmp_path
    ):
        # A market that produced per-campaign TSVs in a PRIOR audit did
        # have campaigns, so a zero now is a regression to justify, not
        # assert. Live: a run that never opened Amazon AE wrote
        # total_active 0 while the previous audit drilled 10 campaigns.
        _setup(
            monkeypatch,
            tmp_path,
            't-hist',
            scope=[
                *_SA_SCOPE,
                {
                    'platform': 'amazon',
                    'country': 'AE',
                    'active_ids': [],
                    'total_active': 0,
                    'total_active_source': 'chip:Live 0',
                },
            ],
            targets={'amazon': ['SA', 'AE']},
        )
        # Prior audit history for THIS store's amazon/ae.
        hist = tmp_path / 'stores' / 'acme' / 'ads' / 'amazon' / 'ae'
        hist.mkdir(parents=True)
        (hist / '100000000001.tsv').write_text('target\tspend\n')
        (hist / '100000000001.searchterms.tsv').write_text('query\tspend\n')
        targets = tmp_path / 'tasks' / 't-hist' / 'AUDIT_TARGETS.json'
        data = json.loads(targets.read_text())
        data['slug'] = 'acme'
        targets.write_text(json.dumps(data))
        acr.reset_progress('t-hist')
        gaps = _gaps(_report(_DRILLED), 't-hist')
        assert any('无在投活动' in g and 'amazon AE' in g for g in gaps), gaps

    def test_empty_combo_with_no_history_passes(self, monkeypatch, tmp_path):
        # A market that never had campaigns is legitimately empty — the
        # history challenge must not fire, or a genuinely-idle marketplace
        # becomes unauditable.
        _setup(
            monkeypatch,
            tmp_path,
            't-nohist',
            scope=[
                *_SA_SCOPE,
                {
                    'platform': 'amazon',
                    'country': 'AU',
                    'active_ids': [],
                    'total_active': 0,
                    'total_active_source': 'chip:Live 0',
                },
            ],
            targets={'amazon': ['SA', 'AU']},
        )
        targets = tmp_path / 'tasks' / 't-nohist' / 'AUDIT_TARGETS.json'
        data = json.loads(targets.read_text())
        data['slug'] = 'acme'
        targets.write_text(json.dumps(data))
        acr.reset_progress('t-nohist')
        gaps = _gaps(_report(_DRILLED), 't-nohist')
        # The emptiness CHALLENGE must not fire. The combo still owes a
        # report section saying it is empty, so that gap is expected.
        assert not any('无在投活动' in g for g in gaps), gaps
        assert not any('exhaustive' in g for g in gaps), gaps

    def test_collapsed_active_count_is_challenged(self, monkeypatch, tmp_path):
        # total_active == len(active_ids) proves internal consistency, not
        # observation — trivially true when both come from the same
        # agent-side derivation. Live: a run declared noon SA 4/4 where a
        # verified audit the same day found 8, because its parser silently
        # dropped TSVs it could not read.
        _setup(
            monkeypatch,
            tmp_path,
            't-collapse',
            scope=[
                *_SA_SCOPE,
                {
                    'platform': 'noon',
                    'country': 'SA',
                    'active_ids': ['C_A1', 'C_A2'],
                    'total_active': 2,
                    'total_active_source': 'chip:Live 2',
                },
            ],
            targets={'amazon': ['SA'], 'noon': ['SA']},
        )
        hist = tmp_path / 'stores' / 'acme' / 'ads' / 'noon' / 'sa'
        hist.mkdir(parents=True)
        for i in range(10):  # 10 campaigns historically, 2 declared now
            (hist / f'C_OLD{i}.tsv').write_text('t\n')
            (hist / f'C_OLD{i}.searchterms.tsv').write_text('q\n')
        tgt = tmp_path / 'tasks' / 't-collapse' / 'AUDIT_TARGETS.json'
        d = json.loads(tgt.read_text())
        d['slug'] = 'acme'
        tgt.write_text(json.dumps(d))
        acr.reset_progress('t-collapse')
        gaps = _gaps(_report(_DRILLED), 't-collapse')
        assert any('不到历史的一半' in g and 'noon SA' in g for g in gaps), gaps

    def test_normal_pause_attrition_is_not_challenged(
        self, monkeypatch, tmp_path
    ):
        # Campaigns get paused, so prior >= current is ordinary. Only a
        # COLLAPSE is suspicious; flagging every shortfall would make the
        # check noise. 5 declared against 6 historical must stay clean.
        _setup(
            monkeypatch,
            tmp_path,
            't-attrition',
            scope=[
                {
                    'platform': 'amazon',
                    'country': 'SA',
                    'active_ids': [f'60000000000{i}' for i in range(5)],
                    'total_active': 5,
                    'total_active_source': 'chip:Live 5',
                }
            ],
            targets={'amazon': ['SA']},
        )
        hist = tmp_path / 'stores' / 'acme' / 'ads' / 'amazon' / 'SA'
        hist.mkdir(parents=True)
        for i in range(6):
            (hist / f'10000000000{i}.tsv').write_text('t\n')
            (hist / f'10000000000{i}.searchterms.tsv').write_text('q\n')
        tgt = tmp_path / 'tasks' / 't-attrition' / 'AUDIT_TARGETS.json'
        d = json.loads(tgt.read_text())
        d['slug'] = 'acme'
        tgt.write_text(json.dumps(d))
        acr.reset_progress('t-attrition')
        gaps = _gaps(_report(_DRILLED), 't-attrition')
        assert not any('不到历史的一半' in g for g in gaps), gaps

    def test_prior_tsvs_count_campaigns_not_files(self, tmp_path):
        # Each campaign leaves TWO files; counting files doubles every
        # total and made the shrink check flag 21-of-22 combos as
        # collapsed.
        d = tmp_path / 'stores' / 'acme' / 'ads' / 'amazon' / 'SA'
        d.mkdir(parents=True)
        (d / '100000000001.tsv').write_text('t\n')
        (d / '100000000001.searchterms.tsv').write_text('q\n')
        (d / '100000000002.tsv').write_text('t\n')
        orig = sc.VIBE_SELLER_DIR
        try:
            sc.VIBE_SELLER_DIR = tmp_path
            assert sc.prior_campaign_tsvs('acme', 'amazon', 'SA') == 2
        finally:
            sc.VIBE_SELLER_DIR = orig

    def test_prior_tsvs_never_borrow_another_stores_history(self, tmp_path):
        # Counting must be scoped to the store's own directory — a glob
        # across stores/* would let one store's history challenge another.
        monkey = tmp_path
        (monkey / 'stores' / 'other' / 'ads' / 'amazon' / 'ae').mkdir(
            parents=True
        )
        (
            monkey / 'stores' / 'other' / 'ads' / 'amazon' / 'ae' / 'x.tsv'
        ).write_text('t\n')
        orig = sc.VIBE_SELLER_DIR
        try:
            sc.VIBE_SELLER_DIR = monkey
            assert sc.prior_campaign_tsvs('acme', 'amazon', 'AE') == 0
            assert sc.prior_campaign_tsvs('other', 'amazon', 'AE') == 1
            assert sc.prior_campaign_tsvs(None, 'amazon', 'AE') == 0
        finally:
            sc.VIBE_SELLER_DIR = orig

    def test_prior_tsvs_ignore_stale_history(self, tmp_path):
        # Without a recency window this count only grows — TSVs are never
        # deleted, so a store audited for a year accumulates every campaign
        # it ever ran and the collapse threshold drifts toward always
        # firing. Only campaigns drilled recently count as evidence that
        # the marketplace currently has campaigns.
        d = tmp_path / 'stores' / 'acme' / 'ads' / 'amazon' / 'SA'
        d.mkdir(parents=True)
        old = 1_600_000_000  # years ago
        for i in range(5):
            f = d / f'10000000000{i}.tsv'
            f.write_text('t\n')
            os.utime(f, (old, old))
        for i in range(5, 7):
            (d / f'10000000000{i}.tsv').write_text('t\n')  # now
        orig = sc.VIBE_SELLER_DIR
        try:
            sc.VIBE_SELLER_DIR = tmp_path
            assert sc.prior_campaign_tsvs('acme', 'amazon', 'SA') == 2
            # the full history is still reachable when asked for
            assert (
                sc.prior_campaign_tsvs(
                    'acme', 'amazon', 'SA', within_days=36500
                )
                == 7
            )
        finally:
            sc.VIBE_SELLER_DIR = orig

    def test_prior_tsvs_match_dir_casing_insensitively(self, tmp_path):
        # Agent-created paths mix casing in the wild (ads/amazon/SA beside
        # ads/amazon/ae). A case-sensitive lookup works on macOS, whose
        # default volume folds case, and silently returns 0 on Linux/WSL —
        # turning the evidence check into a no-op where it also ships.
        d = tmp_path / 'stores' / 'acme' / 'ads' / 'AMAZON' / 'SA'
        d.mkdir(parents=True)
        (d / '100000000001.tsv').write_text('x')
        orig = sc.VIBE_SELLER_DIR
        try:
            sc.VIBE_SELLER_DIR = tmp_path
            assert sc.prior_campaign_tsvs('acme', 'amazon', 'sa') == 1
            assert sc.prior_campaign_tsvs('acme', 'AMAZON', 'SA') == 1
        finally:
            sc.VIBE_SELLER_DIR = orig

    def test_no_targets_file_keeps_prior_behaviour(self, monkeypatch, tmp_path):
        _setup(monkeypatch, tmp_path, 't-none', scope=_SA_SCOPE)
        assert not any(
            '根本没进' in g for g in _gaps(_report(_DRILLED), 't-none')
        )

    def test_execution_summary_owes_no_marketplaces(
        self, monkeypatch, tmp_path
    ):
        # Same skill, same gates — but nothing here claims to be an audit.
        _setup(
            monkeypatch,
            tmp_path,
            't-exec',
            targets={'amazon': ['SA', 'AE'], 'noon': ['SA']},
        )
        gaps = _gaps('Amazon SA 广告优化执行完毕，12 项已应用。', 't-exec')
        assert gaps == [], gaps


@pytest.mark.unit
class TestTargetingLayerHasRows:
    """A targeting table must name targets, not restate the campaign."""

    def test_aggregate_only_table_is_a_gap(self, monkeypatch, tmp_path):
        _setup(monkeypatch, tmp_path, 't-agg', scope=_SA_SCOPE)
        gaps = _gaps(_report(_AGGREGATE_ONLY), 't-agg')
        assert any('[定向层]' in g for g in gaps), gaps

    def test_real_per_keyword_table_passes(self, monkeypatch, tmp_path):
        _setup(monkeypatch, tmp_path, 't-real', scope=_SA_SCOPE)
        gaps = _gaps(_report(_DRILLED), 't-real')
        assert not any('[定向层]' in g for g in gaps), gaps

    def test_recommendation_cell_saying_jian_yi_is_not_a_header(
        self, monkeypatch, tmp_path
    ):
        # A 建议 cell legitimately contains the word ("建议出价 1.30").
        # Reading it as a new header would lose the rows beneath it.
        _setup(monkeypatch, tmp_path, 't-cell', scope=_SA_SCOPE)
        body = _DRILLED.replace('提高至 1.30', '提取为定向词，建议出价 1.30')
        gaps = _gaps(_report(body), 't-cell')
        assert not any('[定向层]' in g for g in gaps), gaps

    def test_aggregate_only_searchterm_table_is_a_gap(
        self, monkeypatch, tmp_path
    ):
        # The targeting layer's evasion has an exact twin: one 汇总 row
        # standing in for the customer queries. Declared live by a run
        # itself — "13 个低花费 Amazon SA 活动的 Search Terms 表用汇总行
        # 代替逐条 top-15". That model self-corrected; a weaker one would
        # ship it, and the search-term layer was only checked for
        # PRESENCE plus a reconciliation line.
        _setup(monkeypatch, tmp_path, 't-st-agg', scope=_SA_SCOPE)
        body = _DRILLED + (
            '| 搜索词 | 来源关键词 | 点击 | 花费 | 建议 |\n'
            '|---|---|---|---|---|\n'
            '| 汇总（15 词） | — | 40 | 8.00 | 详见 TSV |\n'
        )
        gaps = _gaps(_report(body), 't-st-agg')
        assert any('[搜索词层]' in g for g in gaps), gaps

    def test_real_searchterm_rows_pass(self, monkeypatch, tmp_path):
        _setup(monkeypatch, tmp_path, 't-st-ok', scope=_SA_SCOPE)
        body = _DRILLED + (
            '| 搜索词 | 来源关键词 | 点击 | 花费 | 建议 |\n'
            '|---|---|---|---|---|\n'
            '| blue widget large | widget | 9 | 5.00 | 否定（零转化） |\n'
            '| **合计** | — | 9 | 5.00 | — |\n'
        )
        gaps = _gaps(_report(body), 't-st-ok')
        assert not any('[搜索词层]' in g for g in gaps), gaps

    def test_no_data_page_is_exempt(self, monkeypatch, tmp_path):
        _setup(monkeypatch, tmp_path, 't-nodata', scope=_SA_SCOPE)
        body = _AGGREGATE_ONLY.replace(
            '钻取关键词级后再降价', '该活动定向页无数据'
        )
        assert not any(
            '[定向层]' in g for g in _gaps(_report(body), 't-nodata')
        )


@pytest.mark.unit
class TestReconcileContradictionVsIncompleteness:
    """A mismatch is two different failures and only one is forgivable.

    Measured on one live account (bulk export, both layers, same 30d
    window, 17 enabled SP campaigns): ratio min 0.998, median 1.000,
    max 1.000 — 15 of 17 exactly 1.000. noon's captured layers likewise
    top out at 1.000 (median 0.779; its CQ page attributes only part of
    spend). So nothing legitimately exceeds 1.0, and search-term spend
    ABOVE targeting spend is impossible rather than imprecise.
    """

    def _blk(self, tgt, st):
        return (
            '### 600000000001 | acme widgets 004 manual | SP\n'
            '| 关键词 | 出价 | 点击 | 花费 | ROAS | 建议 |\n'
            '|---|---|---|---|---|---|\n'
            '| widget | 1.20 | 6 | 5.00 | 3.5 | 维持 |\n'
            '| widget large | 1.00 | 4 | 3.00 | 4.1 | 提高至 1.30 |\n'
            '| 搜索词 | 来源 | 点击 | 花费 | 建议 |\n'
            '|---|---|---|---|---|\n'
            '| blue widget | widget | 9 | 5.00 | 否定（零转化） |\n'
            f'搜索词对账: 定向花费 USD {tgt} / 点击 10 = '
            f'搜索词花费 USD {st} / 点击 10 (✓)\n'
        )

    def _deny(self, monkeypatch, tmp_path, tid, tgt, st, extra=''):
        _setup(monkeypatch, tmp_path, tid, scope=_SA_SCOPE)
        body = self._blk(tgt, st) + extra
        return acr.check(_report(body), task_id=tid, track=False)

    def test_searchterm_above_targeting_is_a_contradiction(
        self, monkeypatch, tmp_path
    ):
        d = self._deny(monkeypatch, tmp_path, 't-imposs', '480.00', '600.00')
        assert d is not None
        assert any('[对账·不可能]' in g for g in d.contradictions), d.gaps

    def test_two_percent_over_is_tolerated_as_rounding(
        self, monkeypatch, tmp_path
    ):
        # Live max is 1.000; 1.02 is that plus headroom for 2-decimal
        # rounding and currency formatting, so 1.01 must NOT fire.
        d = self._deny(monkeypatch, tmp_path, 't-round', '100.00', '101.00')
        assert not (d and d.contradictions), d.gaps if d else None

    def test_under_attribution_stays_a_soft_gap(self, monkeypatch, tmp_path):
        # Capture missed rows — incompleteness, which the stall may forgive.
        d = self._deny(monkeypatch, tmp_path, 't-under', '100.00', '20.00')
        assert d is not None
        assert not d.contradictions, d.contradictions
        assert any('[对账]' in g for g in d.gaps), d.gaps

    def test_quarantine_marker_resolves_the_contradiction(
        self, monkeypatch, tmp_path
    ):
        # The second legal answer: don't fix the numbers, but tell the
        # reader they cannot be used. This is what keeps a non-stallable
        # check from being a trap.
        d = self._deny(
            monkeypatch,
            tmp_path,
            't-quar',
            '480.00',
            '600.00',
            extra='⚠️ 数据不可信：本活动两层对账矛盾，请勿执行本活动的出价建议\n',
        )
        assert not (d and d.contradictions), d.gaps if d else None

    def test_half_a_disclaimer_does_not_resolve_it(self, monkeypatch, tmp_path):
        # "data is a bit off" without "do not execute" leaves the rows
        # reading as actionable, which is the outcome being prevented.
        d = self._deny(
            monkeypatch,
            tmp_path,
            't-half',
            '480.00',
            '600.00',
            extra='注：本活动数据有偏差，仅供参考\n',
        )
        assert d is not None and d.contradictions, 'must still be flagged'


@pytest.mark.unit
class TestTotalActiveProvenance:
    """``total_active`` must say where it came from; bulk is verified."""

    def test_missing_source_is_a_gap(self, monkeypatch, tmp_path):
        _setup(
            monkeypatch,
            tmp_path,
            't-nosrc',
            scope=[
                {
                    'platform': 'amazon',
                    'country': 'SA',
                    'active_ids': ['600000000001'],
                    'total_active': 1,
                }
            ],
        )
        gaps = _gaps(_report(_DRILLED), 't-nosrc')
        assert any('total_active_source' in g for g in gaps), gaps

    def test_chip_source_accepted_unverified(self, monkeypatch, tmp_path):
        _setup(monkeypatch, tmp_path, 't-chip', scope=_SA_SCOPE)
        gaps = _gaps(_report(_DRILLED), 't-chip')
        assert not any('total_active_source' in g for g in gaps), gaps

    def test_bulk_source_under_the_export_count_is_a_gap(
        self, monkeypatch, tmp_path
    ):
        # The one part of the contract the server VERIFIES: open the
        # named export and count state=enabled campaign rows itself.

        dl = tmp_path / 'downloads' / 'acme'
        dl.mkdir(parents=True)
        wb = Workbook()
        ws = wb.active
        ws.append(['Product', 'Entity', 'Operation', 'State'])
        for _ in range(6):
            ws.append(['SP', 'Campaign', '', 'enabled'])
        ws.append(['SP', 'Campaign', '', 'paused'])
        wb.save(dl / 'export.xlsx')
        monkeypatch.setattr(sc, 'VIBE_SELLER_DIR', tmp_path)
        assert sc.count_enabled_in_export('export.xlsx') == 6
        _setup(
            monkeypatch,
            tmp_path,
            't-bulk',
            scope=[
                {
                    'platform': 'amazon',
                    'country': 'SA',
                    'active_ids': ['600000000001'],
                    'total_active': 1,
                    'total_active_source': 'bulk:export.xlsx',
                }
            ],
        )
        gaps = _gaps(_report(_DRILLED), 't-bulk')
        assert any('数到 6 个' in g for g in gaps), gaps

    def test_one_export_cannot_verify_several_marketplaces(
        self, monkeypatch, tmp_path
    ):
        # A bulk export is PER-MARKETPLACE (one advertiser account, one
        # market), so an export cited by several combos can be
        # authoritative for at most one of them. Observed live: Amazon AE
        # and AU both cited the SA export because neither has a
        # bulk-operations page, and the naive count told each it had
        # "under-declared" against SA's enabled rows.
        dl = tmp_path / 'downloads' / 'acme'
        dl.mkdir(parents=True)
        wb = Workbook()
        ws = wb.active
        ws.append(['Product', 'Entity', 'Operation', 'State'])
        for _ in range(17):
            ws.append(['SP', 'Campaign', '', 'enabled'])
        wb.save(dl / 'sa.xlsx')
        monkeypatch.setattr(sc, 'VIBE_SELLER_DIR', tmp_path)
        shared = 'bulk:sa.xlsx'
        _setup(
            monkeypatch,
            tmp_path,
            't-shared',
            scope=[
                {
                    'platform': 'amazon',
                    'country': 'SA',
                    'active_ids': ['600000000001'],
                    'total_active': 1,
                    'total_active_source': shared,
                },
                {
                    'platform': 'amazon',
                    'country': 'AE',
                    'active_ids': ['600000000002'],
                    'total_active': 1,
                    'total_active_source': shared,
                },
            ],
        )
        gaps = _gaps(_report(_DRILLED), 't-shared')
        # the real problem is named…
        assert any('同时被多个 combo 引用' in g for g in gaps), gaps
        # …and the bogus under-count is NOT
        assert not any('少算' in g for g in gaps), gaps

    def test_declaring_more_than_the_export_is_fine(
        self, monkeypatch, tmp_path
    ):
        # Over-declaring is legitimate: a campaign that spent inside
        # the window and was paused before the export still deserves a
        # drill.

        dl = tmp_path / 'downloads' / 'acme'
        dl.mkdir(parents=True)
        wb = Workbook()
        ws = wb.active
        ws.append(['Product', 'Entity', 'Operation', 'State'])
        ws.append(['SP', 'Campaign', '', 'enabled'])
        wb.save(dl / 'export.xlsx')
        monkeypatch.setattr(sc, 'VIBE_SELLER_DIR', tmp_path)
        _setup(
            monkeypatch,
            tmp_path,
            't-over',
            scope=[
                {
                    'platform': 'amazon',
                    'country': 'SA',
                    'active_ids': ['600000000001'],
                    'total_active': 1,
                    'total_active_source': 'bulk:export.xlsx',
                }
            ],
        )
        gaps = _gaps(_report(_DRILLED), 't-over')
        assert not any('少算' in g for g in gaps), gaps

    def test_unreadable_export_never_blocks(self, monkeypatch, tmp_path):
        # "Could not check" must not become "failed" — a missing download
        # cannot be allowed to block an otherwise sound report.
        monkeypatch.setattr(sc, 'VIBE_SELLER_DIR', tmp_path)
        assert sc.count_enabled_in_export('absent.xlsx') is None
        _setup(
            monkeypatch,
            tmp_path,
            't-gone',
            scope=[
                {
                    'platform': 'amazon',
                    'country': 'SA',
                    'active_ids': ['600000000001'],
                    'total_active': 1,
                    'total_active_source': 'bulk:absent.xlsx',
                }
            ],
        )
        gaps = _gaps(_report(_DRILLED), 't-gone')
        assert not any('少算' in g for g in gaps), gaps


@pytest.mark.unit
class TestProvenanceParsingIsTolerant:
    """Find the token inside prose; reject only self-referential sources.

    The first cut anchored the patterns, and a live run supplied exactly
    what was asked for wrapped in more context than the schema allowed —
    every one was rejected as "no provenance", which is the same failure
    the 搜索词对账 check already learned: answering more fully must not
    read as not answering.
    """

    def test_bulk_token_found_inside_prose(self):
        s = (
            'bulk:acct123:bulk-acct123-20260628-20260728-1700000000000.xlsx '
            '(state=enabled, 行数=21)'
        )
        kind, ref = sc.classify_total_source(s)
        assert kind == 'bulk'
        assert ref == 'bulk-acct123-20260628-20260728-1700000000000.xlsx'

    def test_chip_token_found_inside_prose(self):
        s = (
            'noon ad manager chip:Live 7 at 2026-07-28T18:08 (scrolled '
            'inner container until distinct link count = 7)'
        )
        assert sc.classify_total_source(s) == ('chip', 7)

    def test_list_form_for_a_site_with_no_bulk_export(self):
        # Amazon AE has NO bulk-operations page (404), so there is no file
        # to name; the campaign-list total is the honest source and needs
        # a token of its own.
        s = 'per-campaign detail enumeration (AE Bulk 404; cm 列表 Live=N=8)'
        assert sc.classify_total_source(s) == ('list', 8)
        assert sc.classify_total_source('list:12') == ('list', 12)

    def test_counting_own_tsvs_is_not_provenance(self):
        # The circular case the requirement exists to catch: the agent
        # citing the files it just wrote is not an independent observation.
        s = (
            'per-campaign TSV enumeration at 2026-07-28T14:15:00Z: '
            '4 active campaigns (all with spend); 0 paused'
        )
        assert sc.classify_total_source(s) == (None, None)

    def test_bulk_wins_when_several_tokens_appear(self):
        # Bulk is the only form the server can verify, so prefer it.
        s = 'chip:Live 9 cross-checked against bulk:export.xlsx'
        assert sc.classify_total_source(s) == ('bulk', 'export.xlsx')
